import sys
import json
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

DIRETORIO_BASE = os.path.dirname(os.path.abspath(__file__))
PASTA_ESTOQUES = os.path.join(DIRETORIO_BASE, 'Estoques')
CAMINHO_MODELO = os.path.join(DIRETORIO_BASE, 'Estoque_Modelo.xlsx')

if not os.path.exists(PASTA_ESTOQUES):
    os.makedirs(PASTA_ESTOQUES)

def garantir_aba(wb, nome_aba):
    if nome_aba in wb.sheetnames:
        return wb[nome_aba]
    return wb.create_sheet(title=nome_aba)

def aplicar_estilo_cabecalho(ws, linha, cabecalhos, azul_escuro="1F4E78"):
    fonte_branca = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fill_azul = PatternFill(start_color=azul_escuro, end_color=azul_escuro, fill_type="solid")
    for col_idx, texto in enumerate(cabecalhos, start=1):
        cell = ws.cell(row=linha, column=col_idx, value=texto)
        cell.font = fonte_branca
        cell.fill = fill_azul
        cell.alignment = Alignment(horizontal="center", vertical="center")

def gerar_backup(banco):
    if os.path.exists(CAMINHO_MODELO):
        wb = openpyxl.load_workbook(CAMINHO_MODELO)
    else:
        wb = openpyxl.Workbook()

    # 1. CADASTROS
    ws_cad = garantir_aba(wb, 'Cadastros')
    if ws_cad.max_row < 3 or ws_cad.cell(3, 1).value is None:
        ws_cad.cell(1, 1, "BASE MODELO - CADASTRO DE PRODUTOS E INSUMOS").font = Font(name="Arial", size=12, bold=True, color="1F4E78")
        cabecalhos_cad = ["Código", "Descrição do Item", "NCM", "Categoria", "Unidade", "Est. Mínimo", "Custo Unitário (R$)"]
        aplicar_estilo_cabecalho(ws_cad, 3, cabecalhos_cad)

    for r in range(4, max(ws_cad.max_row + 1, 100)):
        for c in range(1, 8):
            ws_cad.cell(row=r, column=c, value=None)

    for idx, item in enumerate(banco.get('cadastros', []), start=4):
        ws_cad.cell(row=idx, column=1, value=item.get('codigo', ''))
        ws_cad.cell(row=idx, column=2, value=item.get('descricao', ''))
        ws_cad.cell(row=idx, column=3, value=item.get('ncm', '-'))
        ws_cad.cell(row=idx, column=4, value=item.get('categoria', 'Matéria-Prima'))
        ws_cad.cell(row=idx, column=5, value=item.get('unidade', 'KG'))
        ws_cad.cell(row=idx, column=6, value=float(item.get('minimo', 0)))
        ws_cad.cell(row=idx, column=7, value=float(item.get('custo', 0)))

    # 2. MOVIMENTAÇÕES
    ws_mov = garantir_aba(wb, 'Movimentacoes')
    if ws_mov.max_row < 1 or ws_mov.cell(1, 1).value is None:
        cabecalhos_mov = ["Data", "Tipo", "Código", "Produto", "NCM", "Quantidade", "Observação", "Custo Unit. (R$)", "Valor Total (R$)"]
        aplicar_estilo_cabecalho(ws_mov, 1, cabecalhos_mov)

    for r in range(2, max(ws_mov.max_row + 1, 100)):
        for c in range(1, 10):
            ws_mov.cell(row=r, column=c, value=None)

    for idx, mov in enumerate(banco.get('movimentacoes', []), start=2):
        ws_mov.cell(row=idx, column=1, value=mov.get('data', ''))
        ws_mov.cell(row=idx, column=2, value=mov.get('tipo', ''))
        ws_mov.cell(row=idx, column=3, value=mov.get('codigo', ''))
        ws_mov.cell(row=idx, column=4, value=mov.get('produto', ''))
        ws_mov.cell(row=idx, column=5, value=mov.get('ncm', '-'))
        ws_mov.cell(row=idx, column=6, value=float(mov.get('qtd', 0)))
        ws_mov.cell(row=idx, column=7, value=mov.get('obs', ''))
        ws_mov.cell(row=idx, column=8, value=float(mov.get('valorUnit', 0)))
        ws_mov.cell(row=idx, column=9, value=float(mov.get('valorTotal', 0)))

    # 3. ESTRUTURA BOM
    ws_bom = garantir_aba(wb, 'Estrutura_BOM')
    if ws_bom.max_row < 3 or ws_bom.cell(3, 1).value is None:
        ws_bom.cell(1, 1, "BASE MODELO - FICHA TÉCNICA DE PRODUÇÃO (BOM)").font = Font(name="Arial", size=12, bold=True, color="1F4E78")
        cabecalhos_bom = ["Código PA", "Produto PA", "Código MP", "Matéria-Prima", "Qtd Consumida", "Unidade"]
        aplicar_estilo_cabecalho(ws_bom, 3, cabecalhos_bom)

    for r in range(4, max(ws_bom.max_row + 1, 50)):
        for c in range(1, 7):
            ws_bom.cell(row=r, column=c, value=None)

    for idx, bom in enumerate(banco.get('estruturaBOM', []), start=4):
        ws_bom.cell(row=idx, column=1, value=bom.get('codigoPA', ''))
        ws_bom.cell(row=idx, column=2, value=bom.get('produtoPA', ''))
        ws_bom.cell(row=idx, column=3, value=bom.get('codigoMP', ''))
        ws_bom.cell(row=idx, column=4, value=bom.get('materiaPrima', ''))
        ws_bom.cell(row=idx, column=5, value=float(bom.get('qtdConsumida', 0)))
        ws_bom.cell(row=idx, column=6, value=bom.get('unidade', 'KG'))

    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    nome_backup = f"Backup_Estoque_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    caminho_saida = os.path.join(PASTA_ESTOQUES, nome_backup)
    wb.save(caminho_saida)
    wb.close()

    return {"success": True, "arquivo": nome_backup}

def gerar_relatorio_pdf(banco):
    nome_pdf = f"Relatorio_Contabil_Estoque_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    caminho_pdf = os.path.join(PASTA_ESTOQUES, nome_pdf)

    # Configura documento em A4 Paisagem (Landscape)
    doc = SimpleDocTemplate(
        caminho_pdf,
        pagesize=landscape(A4),
        leftMargin=20,
        rightMargin=20,
        topMargin=20,
        bottomMargin=20
    )

    styles = getSampleStyleSheet()
    
    titulo_style = ParagraphStyle(
        'TituloContabil',
        parent=styles['Heading1'],
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#1F4E78'),
        fontName='Helvetica-Bold'
    )
    
    sub_style = ParagraphStyle(
        'SubTituloContabil',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#555555')
    )

    cell_header = ParagraphStyle(
        'CellHeader',
        parent=styles['Normal'],
        fontSize=8,
        leading=10,
        textColor=colors.white,
        fontName='Helvetica-Bold',
        alignment=1
    )

    cell_body = ParagraphStyle(
        'CellBody',
        parent=styles['Normal'],
        fontSize=7,
        leading=9,
        textColor=colors.HexColor('#222222')
    )

    cell_body_center = ParagraphStyle(
        'CellBodyCenter',
        parent=styles['Normal'],
        fontSize=7,
        leading=9,
        alignment=1,
        textColor=colors.HexColor('#222222')
    )

    cell_body_right = ParagraphStyle(
        'CellBodyRight',
        parent=styles['Normal'],
        fontSize=7,
        leading=9,
        alignment=2,
        textColor=colors.HexColor('#222222')
    )

    elements = []

    # Cabeçalho do PDF
    data_geracao = datetime.now().strftime("%d/%m/%Y às %H:%M:%S")
    elements.append(Paragraph("<b>RT ESQUADRIAS</b> — RELATÓRIO CONTÁBIL DE MOVIMENTAÇÃO DE ESTOQUE", titulo_style))
    elements.append(Paragraph(f"Documento auxiliar contábil retido para conferência e balanço | Gerado em: {data_geracao}", sub_style))
    elements.append(Spacer(1, 12))

    # Tabela Contábil
    headers = ["Data", "Tipo Operação", "Código", "Descrição do Produto", "NCM", "Qtd", "V. Unit. (R$)", "V. Total (R$)", "Observação / Ordem"]
    data_table = [[Paragraph(h, cell_header) for h in headers]]

    valor_total_movimentacoes = 0.0

    for m in banco.get('movimentacoes', []):
        qtd = float(m.get('qtd', 0))
        v_unit = float(m.get('valorUnit', 0))
        v_tot = float(m.get('valorTotal', qtd * v_unit))
        valor_total_movimentacoes += v_tot

        linha = [
            Paragraph(str(m.get('data', '-')), cell_body_center),
            Paragraph(str(m.get('tipo', '-')), cell_body_center),
            Paragraph(str(m.get('codigo', '-')), cell_body_center),
            Paragraph(str(m.get('produto', '-')), cell_body),
            Paragraph(str(m.get('ncm', '-')), cell_body_center),
            Paragraph(f"{qtd:.2f}", cell_body_right),
            Paragraph(f"R$ {v_unit:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."), cell_body_right),
            Paragraph(f"R$ {v_tot:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."), cell_body_right),
            Paragraph(str(m.get('obs', '-')), cell_body)
        ]
        data_table.append(linha)

    # Largura total disponível = 842pt - 40pt (margens) = 802pt
    col_widths = [55, 80, 55, 200, 55, 45, 60, 65, 187]

    tabela = Table(data_table, colWidths=col_widths, repeatRows=1)
    tabela.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E78')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CCCCCC')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8F9FA')])
    ]))

    elements.append(tabela)
    elements.append(Spacer(1, 10))

    # Resumo do Rodapé
    total_formatado = f"R$ {valor_total_movimentacoes:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    resumo_style = ParagraphStyle(
        'ResumoTotal',
        parent=styles['Normal'],
        fontSize=10,
        fontName='Helvetica-Bold',
        alignment=2,
        textColor=colors.HexColor('#1F4E78')
    )
    elements.append(Paragraph(f"Total Acumulado em Movimentações: {total_formatado}", resumo_style))

    doc.build(elements)
    return {"success": True, "arquivo": nome_pdf}

if __name__ == '__main__':
    try:
        raw = sys.stdin.read()
        if raw:
            payload = json.loads(raw)
            acao = payload.get('acao', 'backup')
            banco = payload.get('banco', payload)

            if acao == 'relatorio_contador':
                print(json.dumps(gerar_relatorio_pdf(banco)))
            else:
                print(json.dumps(gerar_backup(banco)))
    except Exception as e:
        print(json.dumps({"success": False, "error": str(e)}))
